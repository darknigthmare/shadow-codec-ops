#!/usr/bin/env python3
"""Build an evidence-led CQC workbook without changing game/source files."""
from pathlib import Path
import argparse, csv, hashlib, json, re, shutil, math
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def safe_text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        # Imported data must never become a spreadsheet formula.
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.startswith(("=", "+", "-", "@")):
            value = "'" + value
        return value[:32760]
    return value


def source_rows(path):
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:6000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(text.splitlines(), dialect=dialect))


def index_entries(path):
    text = path.read_text()
    start = text.index("JSON.parse(") + len("JSON.parse(")
    literal, _ = json.JSONDecoder().raw_decode(text[start:])
    rows = json.loads(literal)
    assert isinstance(rows, list)
    return rows


def add_sheet(wb, name, title, note, headers, rows, widths=None):
    ws = wb.create_sheet(name)
    last = get_column_letter(len(headers))
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=19, bold=True, color="F4F8F2")
    ws["A1"].fill = PatternFill("solid", fgColor="142B24")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = ("Reconstruction provisoire — ne remplace pas le classeur d’origine. " if getattr(wb,"_provisional",False) else "") + note
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A2"].font = Font(name="Aptos", size=10, color="3E5149")
    ws.row_dimensions[2].height = 48 if getattr(wb,"_provisional",False) else 36
    for col, header in enumerate(headers, 1):
        c = ws.cell(4, col, header)
        c.fill = PatternFill("solid", fgColor="28513F")
        c.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 30
    for line, row in enumerate(rows, 5):
        maxlines=1
        for col, value in enumerate(row, 1):
            c = ws.cell(line, col, safe_text(value))
            c.font = Font(name="Aptos", size=10, color="203129")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, str) and value.startswith("https://") and "\n" not in value:
                c.hyperlink = value
                c.font = Font(name="Aptos", size=10, color="176746", underline="single")
            width = widths[col-1] if widths and col <= len(widths) else 26
            maxlines=max(maxlines,sum(max(1,math.ceil(len(part)/max(8,width)))for part in str(value).split("\n")))
        ws.row_dimensions[line].height = min(240,max(36,14*maxlines+8))
    if rows:
        tab = Table(displayName="T" + re.sub(r"[^a-zA-Z0-9]", "", name), ref=f"A4:{last}{4+len(rows)}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)
    ws.freeze_panes = "D5" if len(headers) > 4 else "A5"
    ws.sheet_view.showGridLines = False
    for col in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(col)].width = widths[col-1] if widths and col <= len(widths) else 26
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.auto_filter.ref = f"A4:{last}{max(5,4+len(rows))}"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--runtime", type=Path, required=True)
    ap.add_argument("--history", type=Path, required=True)
    ap.add_argument("--checkpoint", type=Path, required=True)
    ap.add_argument("--narrative-root", type=Path)
    ap.add_argument("--descriptor-inventory", type=Path)
    ap.add_argument("--production-proofs", type=Path)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--p42-manifest", type=Path)
    ap.add_argument("--name", default="CQC_SUIVI_CUMULATIF_P42_PREPARATION_V1.xlsx")
    ap.add_argument("--provisional", action="store_true")
    a = ap.parse_args()
    a.output.mkdir(parents=True, exist_ok=True)
    target = a.output / a.name
    assert not target.exists(), "Use a new version; preserve earlier workbook."
    checkpoint = json.loads(a.checkpoint.read_text())
    assert len(checkpoint["entries"]) == 30
    assert sum(bool(x.get("playable")) for x in checkpoint["entries"]) == 24
    sources = [
        ("Profils historiques réconciliés", a.history/"PROFILS_MANQUANTS_RECONCILIES.csv"),
        ("Costumes historiques V3", a.history/"preserved-files/CQC_SPRITES_REPRISE_V3/INVENTAIRE_COSTUMES_V3.csv"),
        ("File historique V3", a.history/"preserved-files/CQC_SPRITES_REPRISE_V3/FILE_GENERATION_V3.csv"),
        ("Checkpoint P41", a.checkpoint),
        ("Index natif P40", a.runtime/"src/cqc-pass40-native-batch-index-v2.js"),
        ("Index natif P41", a.runtime/"src/cqc-pass41-native-batch-index-v1.js"),
    ]
    if a.narrative_root:
        sources.extend([
            ("Scènes historiques v0.56",a.narrative_root/"SCENES_1692_A_PEINDRE_v0.56.csv"),
            ("Récits historiques v0.56",a.narrative_root/"RECITS_282_A_PEINDRE_v0.56.csv"),
        ])
    pins = [{"source": name, "file": path.name, "bytes":path.stat().st_size, "sha256":sha(path.read_bytes())} for name,path in sources]
    profiles, costumes, queue = [source_rows(path) for _,path in sources[:3]]
    native40, native41 = [index_entries(path) for _,path in sources[4:6]]
    scenes,stories=[source_rows(path) for _,path in sources[6:8]]if a.narrative_root else([],[])
    original_profiles_path=a.history/"preserved-files/CQC_SPRITES_REPRISE_V3/PROFILS_MANQUANTS_V3.csv"
    original_profiles=source_rows(original_profiles_path)
    descriptor_pairs=json.loads(a.descriptor_inventory.read_text())["pairs"]if a.descriptor_inventory else[]
    production_evidence=[]
    if a.production_proofs:
        assert checkpoint["production"]["state"]=="READY"
        proofnames=["P41_PRODUCTION_READY_FRESH_ACTUAL_V1.json","P41_PRODUCTION_43_SOURCES_3_ROUTES_STRICT_READBACK_ACTUAL_V1.json","P41_PRODUCTION_11_READONLY_STRICT_ACTUAL_V1.json"]
        for name in proofnames:
            p=a.production_proofs/name;raw=p.read_bytes();d=json.loads(raw)
            assert not re.search(r"/(?:workspace|tmp)/|sediment:\x2f\x2f|file_0{4}|exec-[a-f0-9]{8}-",raw.decode())
            production_evidence.append({"file":p.name,"bytes":len(raw),"sha256":sha(raw),"data":d})
        ready,read43,read11=[x["data"]for x in production_evidence]
        assert ready["deployment"]["id"]==checkpoint["production"]["deployment"] and ready["deployment"]["state"]=="READY"
        assert read43["status"]==read11["status"]=="PASS"
        assert read43["sourceCommit"]==read11["sourceGitCommit"]==checkpoint["finalGitCommit"]
        assert read43["runtimeSourcesVerified"]==43 and read43["canonicalRoutesVerified"]==3 and read11["responseCount"]==11
        assert read43["TLSVerified"]and read11["TLSVerificationEnabled"]and not read43["browserGameplayOrVisualProductionQAClaimed"]and not read11["browserVisualQAClaimed"]
    assert len(native40)==19 and len(native41)==24
    keys = [(r["uid"],r["id"]) for r in native40+native41]
    assert len(set(keys))==43
    batch42 = json.loads(a.p42_manifest.read_text()) if a.p42_manifest else None
    # The manifest shape is deliberately supplied by the coordinating agent.
    planned42 = batch42.get("entries", []) if isinstance(batch42,dict) else []
    fixed42=bool(isinstance(batch42,dict)and batch42.get("schema")=="cqc.pass42.fixed-batch30/1")
    if batch42 is not None:
        assert len(planned42)==30, "A frozen lot is exactly thirty distinct entries."
    wb=Workbook();wb.remove(wb.active);wb._provisional=a.provisional
    now=datetime.now(ZoneInfo("Europe/Paris")).isoformat(timespec="seconds")
    source_observed_paris=datetime.fromisoformat(checkpoint["production"]["lastObservedAtUTC"].replace("Z","+00:00")).astimezone(ZoneInfo("Europe/Paris")).isoformat(timespec="seconds")
    overview=[
        ["Lot P41 — entrées suivies",30,"30 = 24 costumes jouables + 6 entrées de combat en attente"],
        ["P41 — costumes validés localement",24,"Versus24 / Soutien24 / Core10"],
        ["P41 — portraits Hideo",3,"Dossier seulement ; ces portraits ne rendent pas les corps jouables"],
        ["P41 — entrées de combat en attente",6,"1 costume Rétro + 3 corps Hideo + 2 appareils"],
        ["Index natifs P40 + P41",43,"19 précédents conservés + 24 nouveaux ; ne couvre pas toutes les anciennes tenues du jeu"],
        ["Lot P42 — objectifs figés"if fixed42 else"Lot P42 — candidats proposés",len(planned42),"Lot figé ; trente objectifs, pas trente costumes terminés"if fixed42 else"En attente du manifeste coordonné" if not planned42 else "Proposition non figée ; références et absence de doublons à valider"],
        ["Costumes historiques V3",len(costumes),"Historique importé sans déclarer leur état courant validé"],
        ["Profils historiques à réauditer",len(profiles),"Référence, présence et intégration courantes à vérifier séparément"],
        ["Ancienne file de génération V3",len(queue),"Historique ; ne constitue ni une preuve de génération ni une autorisation de réessai"],
        ["Dernier commit source P41 observé",checkpoint["finalGitCommit"],"Source du checkpoint local"],
        ["Vercel — état vérifié",checkpoint["production"]["state"],"Reçus READY + HTTPS 43 sources / 3 routes + 11 réponses indépendantes"if production_evidence else"Observation historique du checkpoint ; pas une nouvelle lecture réseau"],
        ["Vercel — date de cette observation (Paris)",source_observed_paris,"Convertie depuis l'horodatage du reçu source ; pas une nouvelle requête du classeur"],
        ["Objectif fidélité", "1:1 à viser", "Références et incarnations séparées ; aucun certificat universel de reproduction retail pixel par pixel"],
        ["Génération du classeur (Paris)",now,"Les sources ci-dessous ont été lues sans modification"],
    ]
    if a.narrative_root:
        overview.extend([
            ["Scènes citées dans la conversation",len(scenes),"Table historique de scènes à peindre ; état actuel non réaudité"],
            ["Récits cités dans la conversation",len(stories),"Table historique des récits ; état actuel non réaudité"],
        ])
    if fixed42:
        overview.extend([
            ["P42 — répartition des objectifs","10 camouflages / 10 Rétro / 10 Gala","Les incarnations restent distinctes ; trois UID secondaires ont deux familles de costume"],
            ["P42 — correction des doublons","Tiger Stripe exclu ; Water ajouté","Tiger Stripe est déjà la tenue de base. L'intitulé Squares suit le jeu original"],
            ["P42 — validation de référence","En cours","Le plan ne certifie ni génération achevée, ni fidélité intégrale, ni intégration de ces trente costumes"],
        ])
    add_sheet(wb,"Synthèse","CQC — suivi cumulatif des lots de 30","Les états locaux, la disponibilité en production, les portraits et les corps jouables sont suivis séparément.",["Indicateur","Valeur","Périmètre / interprétation"],overview,[43,32,105])
    p41rows=[]
    for i,e in enumerate(checkpoint["entries"],1):
        isplay=bool(e.get("playable"));portrait=bool(e.get("portraitAvailable"))
        status="Validé localement — sources en production vérifiées"if isplay and production_evidence else"Validé localement" if isplay else "Partiel — portrait seulement" if portrait else "En attente — non admis"
        reason="Références du corps complet manquantes" if portrait else "Validation/génération non aboutie ; ne pas marquer terminé" if not isplay else "Références et jeu locaux contrôlés"
        production_state="READY — fichiers publics vérifiés"if production_evidence and isplay else"READY — portrait vérifié, corps non admis"if production_evidence and portrait else"Non admis"if not isplay else"À vérifier séparément"
        p41rows.append([i,e["entryID"],e.get("label",""),e["kind"],e.get("family",""),e.get("incarnation",""),status,"Oui"if isplay else"Non","Oui"if portrait else"Non","Oui"if e.get("modes",{}).get("Versus")else"Non","Oui"if e.get("modes",{}).get("Support")else"Non","Oui"if e.get("modes",{}).get("Core")else"Non",e["status"],reason,production_state])
    add_sheet(wb,"Lot_P41","P41 — trente entrées, état observé","Les six entrées en attente restent identifiées. La colonne d'état brut de l'entrée conserve sa trace d'installation ; l'observation de production actuelle est séparée.",["N°","Clé stable","Entrée","Type","Famille","Incarnation","État vérifié","Corps jouable","Portrait livré","Versus","Soutien","Core","État brut à l'installation","Limite / suite","Observation production actuelle"],p41rows,[6,63,39,15,15,57,43,16,15,12,12,12,55,72,58])
    native_rows=[]
    for lot,entries in [("P40",native40),("P41",native41)]:
        for e in entries:
            p=e.get("provenance",{});rr=e.get("review",{});refs=p.get("sources",[])
            urls="\n".join(x.get("url","") if isinstance(x,dict)else str(x)for x in refs)
            native_rows.append([f"{e['uid']}:{e['id']}",lot,e["uid"],e["id"],e.get("label",""),e.get("family",""),p.get("incarnation",""),"Validé localement — index installé", "1:1 à viser", "Création originale CQC"if p.get("originalDesign")else"Apparence canonique référencée", "Animation/silhouette 2D interprétées ; reproduction retail intégrale non certifiée",urls,rr.get("artAuditSHA256","")])
    add_sheet(wb,"Natifs_indexés","Tenues natives indexées — P40 et P41","Lecture JSON des index installés, sans exécuter leur JavaScript. Inventaire partiel des tenues du jeu.",["Clé stable","Lot","UID","Costume ID","Libellé","Famille","Incarnation","État observé","Cible fidélité","Statut canon / original","Limites vérifiables","Références","Empreinte revue art"],native_rows,[64,10,42,31,32,17,57,43,18,32,75,85,68])
    if descriptor_pairs:
        pairrows=[[r["uid"],r["id"],r.get("family",""),r.get("source",""),"Descripteur statique présent ; composition et sélection actives non réauditées"]for r in descriptor_pairs]
        add_sheet(wb,"Index_statiques","Descripteurs natifs installés — inventaire statique","La présence dans un index n'est pas assimilée à une vérification du corps, des modes ou de toutes les options actives.",["UID","Costume ID","Famille","Fichier index","Portée de l'observation"],pairrows,[48,37,23,65,95])
    historical=[]
    for r in costumes:
        historical.append([r["Identifiant"],r["Personnage"],r["Épisode"],r["Variante"],r["Catégorie"],r["État"],"Non réaudité actuellement",r["Référence prête"],r["Sources"],r["Limites"],r["Pack du PNG"],r["Fichier PNG"]])
    add_sheet(wb,"Historique_costumes","Historique V3 — 416 costumes","Le statut historique est conservé littéralement. Il ne vaut pas preuve de présence, de fidélité ni de jouabilité dans la version actuelle.",["Identifiant source","Personnage","Épisode","Variante","Catégorie","État historique","État courant","Référence prête source","Sources","Limites source","Pack historique","Fichier historique"],historical,[65,30,26,52,24,27,32,23,85,100,34,72])
    profrows=[]
    for r in profiles:
        profrows.append([r["uid"],r["name"],r["episode"],r["status"],"Non réaudité actuellement",r["verified_elements"],r["fidelity_limitations"],r["reference_urls"],r["reason"],r["canonical_fidelity_certified"]])
    add_sheet(wb,"Profils_à_réauditer","Profils historiques — éléments manquants","Les 45 lignes d'origine sont préservées. L'existence d'une référence ou d'un candidat ne permet pas de déclarer le profil jouable.",["UID","Personnage","Épisode","État historique","État courant","Éléments établis","Limites de fidélité","Références","Motif source","Certifié dans la source"],profrows,[52,31,25,26,32,85,100,85,100,24])
    origprows=[[r[k]for k in ["Identifiant","Personnage","Épisode","État","Référence prête","Motif"]]+["État historique ; non réaudité actuellement"]for r in original_profiles]
    add_sheet(wb,"Profils_V3_originaux","Profils V3 — le tableau original préservé","Ces six colonnes sont celles du CSV de l'archive ; le réaudit local est suivi dans l'onglet distinct Profils_à_réauditer.",["Identifiant","Personnage","Épisode","État source","Référence prête source","Motif source","Qualification courante"],origprows,[55,34,31,26,24,103,56])
    qrows=[[r["Identifiant"],r["Personnage"],r["Épisode"],r["Variante"],r["Type"],r["Nouvellement prêt V3"],r["Nombre de références"],r["Sources"],"Ancienne file ; production courante non vérifiée"]for r in queue]
    add_sheet(wb,"File_historique","Ancienne file V3 — 61 costumes","Ce tableau reste un historique de travail, séparé du lot courant de trente.",["Identifiant source","Personnage","Épisode","Variante","Type","Prêt V3 (source)","Références (source)","Sources","Qualification courante"],qrows,[67,31,30,55,15,24,24,95,57])
    if a.narrative_root:
        scenerows=[[r[k]for k in ["sceneKey","uid","fighter","game","phase","index","title","place","priority","status","portraitReference"]]+["État historique ; rendu actuel non réaudité"]for r in scenes]
        add_sheet(wb,"Scènes_historiques","Scènes v0.56 — 1 692 lignes citées dans la conversation","Source retrouvée localement dans l'archive ; aucun état courant n'est déduit de cette ancienne file.",["Clé de scène","UID","Combattant","Épisode","Phase","Index","Titre","Lieu","Priorité source","État historique","Portrait de référence source","État courant"],scenerows,[62,43,31,25,15,10,49,42,18,26,84,53])
        storyrows=[[r[k]for k in ["uid","fighter","game","group","fullSceneCount","missingSceneCount","priority","priorityOrder","portraitReference"]]+["État historique ; contenu actuel non réaudité"]for r in stories]
        add_sheet(wb,"Récits_historiques","Récits v0.56 — 282 lignes citées dans la conversation","Les nombres de scènes complètes/manquantes sont ceux de la source historique, pas une nouvelle inspection du jeu.",["UID","Combattant","Épisode","Groupe","Scènes complètes source","Scènes manquantes source","Priorité source","Ordre source","Portrait de référence source","État courant"],storyrows,[47,32,26,44,23,23,18,17,87,57])
    p42rows=[]
    for i,e in enumerate(planned42,1):
        refs=e.get("reference_urls",e.get("references",[]))
        if isinstance(refs,list):refs="\n".join((x.get("url","")+" — "+x.get("status","")+" — "+x.get("scope",""))if isinstance(x,dict)else str(x)for x in refs)
        limitation = "Création originale CQC ; identité et incarnation référencées, sans costume canonique commercial correspondant" if e.get("originalDesign") else "Base actuelle à comparer" if e.get("existingNativeBaseUniformIdentityMustBeCompared") else e.get("limitation",e.get("notes",""))
        p42rows.append([i,e.get("entryID",e.get("key",e.get("id",e.get("uid","")))),e.get("label",e.get("name","")),e.get("uid",""),e.get("family",e.get("kind",e.get("type",""))),e.get("incarnation",e.get("episode","")),"Références à vérifier","1:1 à viser",refs,limitation,e.get("status","En préparation"),"Non — candidat","Non"])
    p42sheet=add_sheet(wb,"Lot_P42","P42 — trente objectifs figés"if fixed42 else"P42 — trente emplacements de travail","Lot figé : trente objectifs distincts. La référence, la génération, la validation artistique et l'intégration restent suivies séparément ; aucun objectif n'est automatiquement terminé."if fixed42 else"Aucune ligne fictive : les trente entrées seront ajoutées depuis la proposition coordonnée. Un état planifié ne compte pas comme terminé."if not planned42 else"Proposition non figée : trente candidats ; aucun n'est déclaré généré ou jouable du seul fait de figurer dans ce tableau.",["N°","Clé stable","Entrée","UID","Famille","Incarnation","État réel","Cible fidélité","Références et état de lecture","Limite / suivi","État source du manifeste","Prêt à générer","Jouable localement"],p42rows,[6,65,38,45,20,57,30,21,90,58,75,25,25])
    if planned42:
        assert len({r[1]for r in p42rows})==30
        choices=DataValidation(type="list",formula1='"Références à vérifier,En préparation,Génération partielle,À contrôler,Validé localement,En attente"',allow_blank=False)
        choices.errorTitle="État du lot";choices.error="Choisir un état de suivi.";choices.showErrorMessage=True;p42sheet.add_data_validation(choices);choices.add("G5:G34")
    refsrows=[[x["source"],x["file"],x["bytes"],x["sha256"],"Lecture locale — origine préservée"]for x in pins]
    refsrows.append(["GitHub — revue du code","PR3","","","https://github.com/darknigthmare/shadow-codec-ops/pull/3"])
    refsrows.extend([
        ["Conversation — scènes et récits","Partage1","","","https://chatgpt.com/share/6abdc55d-6938-83eb-a5e8-682b7717c7dc"],
        ["Conversation — reprise des sprites","Partage2","","","https://chatgpt.com/share/6abdc758-1300-83eb-bc99-ad222967908c"],
        ["Conversation — inventaires de costumes","Partage3","","","https://chatgpt.com/share/6abdc82b-b644-83ed-9064-c5a0d0928a3c"],
    ])
    refsrows.append(["Profils V3 originaux",original_profiles_path.name,original_profiles_path.stat().st_size,sha(original_profiles_path.read_bytes()),"Six colonnes d'origine conservées séparément du réaudit"])
    if a.descriptor_inventory:refsrows.append(["Inventaire statique des index",a.descriptor_inventory.name,a.descriptor_inventory.stat().st_size,sha(a.descriptor_inventory.read_bytes()),"Présence statique seulement ; pas une validation de composition active"])
    if a.p42_manifest:refsrows.append(["Manifeste P42 figé"if fixed42 else"Proposition P42",a.p42_manifest.name,a.p42_manifest.stat().st_size,sha(a.p42_manifest.read_bytes()),"Trente objectifs ; ne prouve pas leur achèvement"])
    for p in production_evidence:refsrows.append(["Preuve production P41",p["file"],p["bytes"],p["sha256"],"READY et vérification HTTPS ; aucune validation visuelle en navigateur distant revendiquée"])
    add_sheet(wb,"Sources","Sources et empreintes","Les chemins privés et identifiants d'outils ne sont pas incorporés au classeur. Les URL de référence des médias restent dans leurs tableaux.",["Source","Fichier","Octets","SHA256","Qualification / lien"],refsrows,[39,62,17,69,79])
    conversation_rows=[
        ["https://chatgpt.com/share/6abdc82b-b644-83ed-9064-c5a0d0928a3c","assistant — final","CQC_FILE_GENERATION_V3.csv","liste des 61 variantes prêtes",sources[2][1].name,pins[2]["sha256"],"Noms et compte concordants ; ancien téléchargement sandbox non récupéré"],
        ["https://chatgpt.com/share/6abdc82b-b644-83ed-9064-c5a0d0928a3c","assistant — final","CQC_INVENTAIRE_COSTUMES_V3.csv","état des 416 variantes répertoriées",sources[1][1].name,pins[1]["sha256"],"Noms et compte concordants ; ancien téléchargement sandbox non récupéré"],
        ["https://chatgpt.com/share/6abdc82b-b644-83ed-9064-c5a0d0928a3c","assistant — final","CQC_PROFILS_MANQUANTS_V3.csv","les 45 profils encore sans sprite de base", "PROFILS_MANQUANTS_V3.csv",sha(original_profiles_path.read_bytes()),"Version originale préservée ; tableau réconcilié actuel importé séparément"],
    ]
    if a.narrative_root:
        conversation_rows.extend([
            ["https://chatgpt.com/share/6abdc55d-6938-83eb-a5e8-682b7717c7dc","assistant — final","SCENES_1692_A_PEINDRE_v0.56.csv","CSV — 1 692 scènes",sources[6][1].name,pins[6]["sha256"],"Noms et compte concordants ; état actuel non inféré"],
            ["https://chatgpt.com/share/6abdc55d-6938-83eb-a5e8-682b7717c7dc","assistant — final","RECITS_282_A_PEINDRE_v0.56.csv","CSV — 282 récits",sources[7][1].name,pins[7]["sha256"],"Noms et compte concordants ; état actuel non inféré"],
        ])
    add_sheet(wb,"Sources_conversations","Les tableaux cités dans les conversations","Extraits de messages assistant final lus dans les partages publics ; seules les descriptions de fichiers sont citées, sans données de session ni chemins privés.",["Conversation source","Rôle du message","Fichier cité","Description historique exacte","Source locale liée","SHA256 source locale","Portée de la correspondance"],conversation_rows,[85,24,61,58,61,69,94])
    if production_evidence:
        observed=datetime.fromisoformat(checkpoint["production"]["lastObservedAtUTC"].replace("Z","+00:00")).astimezone(ZoneInfo("Europe/Paris")).isoformat(timespec="seconds")
        prodrows=[
            ["Déploiement",checkpoint["production"]["deployment"],"READY",production_evidence[0]["sha256"]],
            ["Commit servi",checkpoint["finalGitCommit"],"Source exacte vérifiée",production_evidence[0]["file"]],
            ["Alias canonique","https://shadow-codec-ops.vercel.app","Vérifié",observed],
            ["Sources et routes", "43 sources + 3 routes", "HTTPS / octets / SHA256 vérifiés",production_evidence[1]["sha256"]],
            ["Contrôle indépendant","11 réponses", "HTTPS / octets / SHA256 vérifiés",production_evidence[2]["sha256"]],
            ["QA de jeu distante","Non revendiquée","Les tests de jeu restent les tests locaux documentés","Pas de nouvelle certification visuelle de production"],
        ]
        add_sheet(wb,"Production_P41","Production P41 — READY et fichiers vérifiés","Les reçus sources sont conservés dans le dossier production_preuves du livrable. Date d'observation affichée en heure de Paris.",["Contrôle","Valeur","Résultat","Preuve / empreinte / date"],prodrows,[35,74,58,82])
    legend=[
        ["Validé localement","Intégré et vérifié dans la version locale ; disponibilité publique à confirmer séparément"],
        ["Partiel — portrait seulement","Portrait livré ; aucun corps de combat jouable ajouté"],
        ["En attente — non admis","Aucun résultat de combat admis pour cette entrée"],
        ["Non réaudité actuellement","État historique conservé ; vérification courante encore nécessaire"],
        ["1:1 à viser","Objectif de fidélité ; ne constitue pas une certification littérale de tout mouvement et détail"],
        ["Création originale CQC","Variante développée pour le jeu, sans prétendre à un costume canonique attesté"],
    ]
    add_sheet(wb,"Légende","Lire les états du suivi","Les personnages à des époques différentes restent des fiches distinctes.",["État / notion","Définition"],legend,[43,125])
    wb.active=0
    wb.properties.creator="CQC project"
    wb.properties.title="CQC — suivi cumulatif, références et lots de trente"
    wb.properties.subject="Suivi de fidélité, génération, intégration et limites réelles"
    wb.save(target)
    # Independent readback: every imported history row and all thirty P41 entries.
    check=load_workbook(target,read_only=True,data_only=False)
    expected={"Synthèse":len(overview),"Lot_P41":30,"Natifs_indexés":43,"Historique_costumes":416,"Profils_à_réauditer":45,"Profils_V3_originaux":45,"File_historique":61,"Lot_P42":len(planned42),"Sources":len(refsrows),"Sources_conversations":len(conversation_rows),"Légende":len(legend)}
    if a.narrative_root:expected.update({"Scènes_historiques":1692,"Récits_historiques":282})
    if descriptor_pairs:expected.update({"Index_statiques":len(descriptor_pairs)})
    if production_evidence:expected.update({"Production_P41":6})
    for name,count in expected.items():
        assert check[name].max_row==max(4,4+count),(name,check[name].max_row,count)
    values=list(check["Lot_P41"].iter_rows(min_row=5,values_only=True))
    assert len(values)==30 and sum(r[7]=="Oui"for r in values)==24 and sum(r[8]=="Oui"for r in values)==3
    imported=list(check["Historique_costumes"].iter_rows(min_row=5,values_only=True))
    assert [r[0]for r in imported]==[r["Identifiant"]for r in costumes]
    assert [r[5]for r in imported]==[r["État"]for r in costumes]
    assert len({r[0]for r in imported})==len(imported)
    assert all(c.data_type!="f" for ws in check for row in ws for c in row),"No imported formula execution"
    check.close()
    snapshot=a.output/"sources_csv";snapshot.mkdir(exist_ok=True)
    for _,path in [row for row in sources if row[1].suffix.lower()==".csv"]:
        dest=snapshot/path.name
        if dest.exists():assert dest.read_bytes()==path.read_bytes()
        else:shutil.copyfile(path,dest)
    if production_evidence:
        pd=a.output/"production_preuves";pd.mkdir(exist_ok=True)
        for proof in production_evidence:
            source=a.production_proofs/proof["file"];dest=pd/proof["file"]
            if dest.exists():assert dest.read_bytes()==source.read_bytes()
            else:shutil.copyfile(source,dest)
    report={"schema":"cqc.excel.cumulative-evidence/1","createdAtParis":now,"provisionalReconstruction":a.provisional,"originalConversationWorkbookRecovered":False,"workbook":{"file":target.name,"bytes":target.stat().st_size,"sha256":sha(target.read_bytes())},"sources":pins,"sheets":expected,"checks":{"P41Entries":30,"P41PlayableCostumes":24,"HideoPortraitsOnly":3,"P41CombatEntriesHeld":6,"NativeIndexedCostumes":43,"HistoricalCostumeRowsPreserved":416,"HistoricalProfileRowsPreserved":45,"HistoricalGenerationQueueRowsPreserved":61,"CurrentHistoricalStatesNotInferred":True,"ImportedFormulas":0,"SourceOrGameFilesModified":False,"RemoteWorkbookModified":False},"discovery":{"localWorkbooksFound":0,"connectedExcelSessionsFound":0,"relatedGoogleDriveWorkbookFound":False},"runtimeGitCommit":checkpoint["runtimeGitCommit"],"sourceGitCommit":checkpoint["finalGitCommit"],"productionStateSourceDate":checkpoint["production"]["lastObservedAtUTC"],"productionStateNotRefetched":checkpoint["production"]["state"]}
    report["p42FixedManifest"]={"fixed":fixed42,"sha256":sha(a.p42_manifest.read_bytes())if a.p42_manifest else None,"entries":len(planned42),"markedPlayable":0}
    report["productionEvidence"]=[{k:v for k,v in p.items()if k!="data"}for p in production_evidence]
    p=a.output/(target.stem+"_CONTROLES.json");assert not p.exists();p.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"workbook":str(target),"sha256":report["workbook"]["sha256"],"bytes":target.stat().st_size,"sheets":expected,"checks":"PASS","report":str(p)},ensure_ascii=False))


if __name__=="__main__":
    main()
